import csv
import hashlib
import io
import json
import re
from pathlib import Path
from .database import dump, now, uid
from .analysis import coverage

MIMES = {
    "json": "application/json",
    "csv": "text/csv",
    "md": "text/markdown",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def cell(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return dump(value)
    return str(value)


def safe_cell(value):
    value = cell(value)
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def save_artifact(con, root, analysis_id, name, mime, content):
    identity = uid()
    name = (
        re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", Path(name.replace("\\", "/")).name).strip(
            " ."
        )[:180]
        or "artifact.bin"
    )
    if name.split(".")[0].upper() in {
        "CON",
        "PRN",
        "AUX",
        "NUL",
        *[f"COM{i}" for i in range(1, 10)],
        *[f"LPT{i}" for i in range(1, 10)],
    }:
        name = "_" + name
    relative = f"artifacts/{identity}/{name}"
    destination = root / relative
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(content)
    con.execute(
        "INSERT INTO artifacts VALUES(?,?,?,?,?,?,?,?)",
        (
            identity,
            analysis_id,
            name,
            mime,
            relative,
            hashlib.sha256(content).hexdigest(),
            len(content),
            now(),
        ),
    )
    return {
        "id": identity,
        "name": name,
        "size": len(content),
        "sha256": hashlib.sha256(content).hexdigest(),
        "download_url": f"/api/v1/artifacts/{identity}/download",
    }


def export_analysis(con, root, analysis_id, format):
    analysis = dict(
        con.execute("SELECT * FROM analyses WHERE id=?", (analysis_id,)).fetchone()
    )
    results = [
        {
            "id": r["id"],
            "payload": json.loads(r["payload"]),
            "evidence_ids": json.loads(r["evidence"]),
        }
        for r in con.execute(
            "SELECT * FROM results WHERE analysis_id=? ORDER BY created,id",
            (analysis_id,),
        )
    ]
    progress = coverage(con, analysis_id)
    evidence = [
        json.loads(r[0])
        for r in con.execute(
            "SELECT snapshot FROM analysis_inputs WHERE analysis_id=? ORDER BY ordinal",
            (analysis_id,),
        )
    ]
    metadata = {
        "name": analysis["name"],
        "goal": analysis["goal"],
        "state": analysis["state"],
        "coverage": progress,
    }
    flat = [
        {"result_id": r["id"], **r["payload"], "evidence_ids": r["evidence_ids"]}
        for r in results
    ]
    fields = list(dict.fromkeys(key for r in flat for key in r)) or [
        "result_id",
        "evidence_ids",
    ]
    if format == "json":
        content = json.dumps(
            {**metadata, "results": results, "evidence": evidence},
            ensure_ascii=False,
            indent=2,
        ).encode("utf-8")
    elif format == "csv":
        out = io.StringIO(newline="")
        writer = csv.DictWriter(out, fieldnames=fields)
        writer.writeheader()
        writer.writerows({k: safe_cell(v) for k, v in row.items()} for row in flat)
        content = out.getvalue().encode("utf-8-sig")
    elif format == "xlsx":
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.title = "Results"
        sheet.append(fields)
        for row in flat:
            sheet.append([safe_cell(row.get(k)) for k in fields])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        summary = book.create_sheet("Coverage")
        for key, value in metadata.items():
            summary.append([key, safe_cell(value)])
        sources = book.create_sheet("Evidence")
        sources.append(
            ["record_id", "platform", "kind", "source_url", "revision", "payload"]
        )
        for r in evidence:
            sources.append(
                [
                    safe_cell(r.get(k))
                    for k in (
                        "id",
                        "platform",
                        "kind",
                        "source_url",
                        "revision",
                        "payload",
                    )
                ]
            )
        for page in book:
            for column in page.columns:
                page.column_dimensions[column[0].column_letter].width = 30
        out = io.BytesIO()
        book.save(out)
        content = out.getvalue()
    else:
        lines = [
            f"# {analysis['name']}",
            "",
            analysis["goal"],
            "",
            f"State: {analysis['state']}",
            f"Coverage: {dump(progress)}",
            "",
        ]
        for r in results:
            lines += [
                f"## {r['id']}",
                *(f"{k}: {cell(v)}" for k, v in r["payload"].items()),
                "Evidence: " + ", ".join(r["evidence_ids"]),
                "",
            ]
        lines += [
            "## Source snapshots",
            *[
                f"{r['id']} | {r['source_url']} | {dump(r['payload'])}"
                for r in evidence
            ],
        ]
        if format == "md":
            content = "\n".join(lines).encode("utf-8")
        else:
            from docx import Document

            document = Document()
            for line in lines:
                if line.startswith("# "):
                    document.add_heading(line[2:], 0)
                elif line.startswith("## "):
                    document.add_heading(line[3:], 1)
                else:
                    document.add_paragraph(line)
            out = io.BytesIO()
            document.save(out)
            content = out.getvalue()
    return save_artifact(
        con,
        root,
        analysis_id,
        f"analysis-{analysis_id}.{format}",
        MIMES[format],
        content,
    )
